"""NSE ASM (Additional Surveillance Measure) ingestion, from dated SURV circulars -- not the live
`/api/reportASM` snapshot, which carries no entry/exit dates and no real history (see
`docs/phase4_asm_gsm_sourcing.md`, "30-second check").

Each periodic ASM circular's Annexure workbook already carries the exact bitemporal shape this
project needs, natively: the circular's own publication date (`cirDate`) is knowledge_date; the
"...w.e.f. DATE" phrase inside each section title is event_date. These are two different, real
dates and must never be collapsed into one (a circular published Jan 6 can announce a change
effective Jan 8 -- see the T1_SURV66001/T2_SURV66016 reconciliation in the phase doc).

Verified (this session, real circulars spanning 2019-2026) that the Annexure sheet is NOT a flat
table: it is a sequence of independently-titled sub-sections ("List of securities shortlisted in
Long - Term ASM Framework Stage - I w.e.f. ...", "...to move from Stage - I to Stage - II w.e.f.
...", "...to be excluded from ASM Framework w.e.f. ...") each with its own header row and a `Nil`
placeholder when empty. A naive single-table parse looks like it disagrees with the Consolidated
sheet; a section-aware parse reconciles exactly (this was the ten-minute check that decided this
design -- see `docs/phase4_asm_gsm_sourcing.md`).

Stage is always taken from the section title text, never from the live API's `gsmStage`-style
field (ASM's live snapshot has no such bug, but the principle -- stage comes from the dated
source document, not a live derived field -- is the same one that ruled out `gsmStage` for GSM).

Explicitly out of scope, not silently dropped: sections whose title mentions the Insolvency and
Bankruptcy Code (IBC) carve-out ("...ASM for Companies relating to the Insolvency Resolution
Process...", present in the 2019-2021 era's "Annexure I-B" / "Consolidated - ASM (IBC)" sheets).
This is a real, different sub-mechanism (IBC-triggered placement, not the numbered Stage I-IV
surveillance criteria) that this build does not parse; every such section title is still counted
and returned in `unparsed_titles` so it is visible, not silently missing. ESM (Enhanced
Surveillance Measure) is a further, entirely separate mechanism this module does not touch at all
-- see CLAUDE.md.

Consolidated sheets are deliberately not ingested as facts here: they are a periodic snapshot of
current membership, not a stream of dated events, and ingesting both the deltas and a snapshot of
their cumulative effect would double-record the same fact two different ways. The delta sections
parsed here are, by the reconciliation check, already a complete and authoritative record.
"""
from __future__ import annotations
import re
import time
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Callable

import openpyxl
import requests

from ...bitemporal.store import StoreValidationError, write_facts

ASM_LT = "ASM_LT"
ASM_ST = "ASM_ST"
ENTRY = "ENTRY"
EXIT = "EXIT"
STAGE_CHANGE = "STAGE_CHANGE"

# Longest-alternative-first purely for readability; the trailing \b makes ordering irrelevant to
# correctness (a partial match of "I" against "IV" fails the boundary check and backtracks).
_STAGE_ALT = r"(III|II|IV|I)"

ENTRY_TITLE_RE = re.compile(
    rf"shortlisted in (Long|Short)\s*-\s*Term ASM Framework Stage\s*-\s*{_STAGE_ALT}\b", re.I)
TRANSITION_TITLE_RE = re.compile(
    rf"shortlisted to move from (Long|Short)\s*-\s*Term ASM Framework Stage\s*-\s*{_STAGE_ALT}"
    rf"\s*to\s*Stage\s*-\s*{_STAGE_ALT}\b", re.I)
# The 2023-era ST-ASM annexure phrases its exclusion title as "excluded from Short - Term ASM
# Framework" -- other eras (2019/2021 LT, 2025/2026 ST) say the bare "excluded from ASM
# Framework". Both real forms must match (P4-001, docs/DEFECT_REGISTER.md).
EXCLUSION_TITLE_RE = re.compile(r"excluded from (?:(Long|Short)\s*-\s*Term\s+)?ASM Framework\b", re.I)
TITLE_START_RE = re.compile(r"^\s*list of securities\b", re.I)
WEF_DATE_RE = re.compile(r"w\.e\.f\.?\s*([A-Za-z]+\s+\d{1,2},?\s*\d{4})", re.I)
_MONTH_DATE_RE = re.compile(r"([A-Za-z]+)\s+(\d{1,2}),?\s*(\d{4})")

def to_iso_date(raw: str) -> str:
    m = _MONTH_DATE_RE.search(raw)
    if not m:
        raise ValueError(f"Cannot parse a month/day/year date out of {raw!r}.")
    month_name, day, year = m.groups()
    text = f"{month_name} {day} {year}"
    # Circulars roughly SURV52458-58318 (mid-2022 to mid-2023) spell the "w.e.f." date with an
    # abbreviated, comma-less month ("Apr 03 2023") instead of the full-name form used everywhere
    # else ("April 03, 2023") -- confirmed against 232 distinct real failing date strings, all one
    # shape (P4-006, docs/DEFECT_REGISTER.md). Try full name first (the common case), then
    # abbreviated.
    for fmt in ("%B %d %Y", "%b %d %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Cannot parse a month/day/year date out of {raw!r} (tried full and abbreviated month names).")

def mechanism_from_subject(subject: str) -> str:
    return ASM_ST if re.search(r"short\s*-?\s*term", subject or "", re.I) else ASM_LT

def _normalize(s: str) -> str:
    return re.sub(r"[^a-z]", "", (s or "").lower())

# P4-013 (docs/DEFECT_REGISTER.md): `is_periodic_asm_subject` was originally an ALLOWLIST
# (require "applicab"+"surve"+"measure" stems, with narrow fallbacks for specific observed
# typos/omissions -- P4-007, P4-012). Two independent real circulars each dropped a DIFFERENT
# required word entirely from an otherwise-genuine periodic subject (P4-012: "Surveillance"
# dropped from SURV63984; then "Measure" dropped from SURV49492) -- each one silently and
# invisibly skipped an entire real circular, only found by tracing one affected symbol's history
# by hand. Two distinct misses of the same shape means a third is likely, and an allowlist false
# negative is silent -- discoverable only by exhaustively tracing individual symbols across dozens
# of circulars, exactly what found these two. A denylist false positive is the opposite: it gets
# attempted, fails to parse (no matching Annexure, or a zip/PDF shape nothing recognizes), and
# shows up immediately and loudly as a circulars_failed entry with the reason attached. That
# asymmetry is why this function is inverted: attempt every SURV circular UNLESS it matches a
# known-irrelevant category below. The categories were built from a full survey of all 1,756
# distinct real subjects in the cached 2019-2026 circular index (not guessed) -- see
# `docs/phase4_asm_gsm_sourcing.md` for the dry-run comparison against the old allowlist run
# before this design was ever used against the network.
_DENY_PATTERNS: tuple[tuple[str, re.Pattern], ...] = (
    ("GSM", re.compile(r"graded\s+surveillance|\bgsm\b", re.I)),
    ("ESM", re.compile(r"enhanced\s+surveillance|\besm\b|\benhnaced\b", re.I)),
    ("IBC", re.compile(r"insolvency|\bibc\b", re.I)),
    ("ICA", re.compile(r"inter\s*creditor|\bica\b", re.I)),
    # \w* tolerates the real observed typo "Encumberance" (extra 'e') as well as the correct spelling.
    ("Encumbrance", re.compile(r"encumb\w*ance", re.I)),
    ("Pledge", re.compile(r"pledge", re.I)),
    ("Deep-OTM", re.compile(r"deep.{0,20}(out.?of.?the.?money|\botm\b)", re.I)),
    ("Trade-for-Trade", re.compile(r"trade\s+for\s+trade|\btft\b", re.I)),
    ("Persistent-Noise", re.compile(r"persistent\s+noise|order\s+based\s+surveillance", re.I)),
    ("Low-NonPromoter-Holding", re.compile(r"non.?promoter\s+holding", re.I)),
    ("Policy-Framework-Update", re.compile(
        r"framework\s*-?\s*update|usage\s+of\s+financial\s+parameters|framework\s+on\s+equity\s+derivatives|"
        r"framework\s+for\s+sme\s+segment|framework\s+on\s+public\s+sector|\bPSU\b|"
        r"public\s+sector\s+undertaking|enhancing\s+trading\s+convenience|"
        r"framework\s+on\s+material\s+price\s+movement|rumour\s+verification|"
        r"margin\s+for\s+non-?F&?O\s+stocks|increase\s+in\s+margin", re.I)),
    # Deliberately broad and word-fragment-tolerant: real enforcement-circular titles are manually
    # typed one-offs, and typos on THESE words specifically have already been observed twice
    # ("Confimatory" for Confirmatory, "Coriggendum" for Corrigendum) -- \w* / doubled-letter
    # tolerance here is not guessing, it is the same lesson from is_periodic_asm_subject's old
    # allowlist applied to the deny side, where a gap is lower-stakes (noise through, not data
    # loss) but still worth closing before a real run.
    ("Individual-Order-Case", re.compile(
        r"\bSEBI\b|\bSATs?\s?Orders?\b|\bNCLT\b|conf\w*matory\s+orders?|final\s+orders?|interim\s+orders?|"
        r"corr?ig+endum|addendum|directions?\b|hon'?ble|in\s+the\s+matter\s+of|"
        r"disposal\s+of\s+representation|\bPAN\s+in\s+the\s+matter|"
        r"\border\s+(?:in\s+respect\s+of|of)\b", re.I)),
    ("Price-Bands", re.compile(r"price\s+bands?|dynamic\s+price\s+band", re.I)),
    ("Algo-Trading", re.compile(r"\balgo\b|algorithmic\s+trading", re.I)),
    ("Order-to-Trade-Ratio", re.compile(r"order.?to.?trade\s+ratio|\bOTR\b", re.I)),
    ("Surveillance-Admin", re.compile(
        r"surveillance\s+(obligation|indicator|dashboard|action\s+w\.?r\.?t|penalty)|"
        r"consolidated\s+(circular|penalty)(\s+structure)?\s+for\s+surveillance|"
        r"surveillance\s+and\s+investigation\s+consolidated", re.I)),
    ("Standardization-SOP", re.compile(r"standardi[sz]ation|standard\s+operating\s+procedure", re.I)),
    ("USDINR-FX", re.compile(r"USDINR", re.I)),
    ("Withholding-Payout", re.compile(r"withholding\s+payout", re.I)),
    ("Trade-Cancellation", re.compile(r"reversal\s+trade|\bRTCM\b|\bSGTCM\b|same\s+group\s+trade\s+cancellation", re.I)),
    ("Position-Limits", re.compile(r"position\s+limits?|intraday\s+position|\bMWPL\b", re.I)),
    ("ETF", re.compile(r"exchange\s+traded\s+funds?|\bETFs?\b", re.I)),
    ("Broker-Confidence-Measures", re.compile(r"instil\s+confidence|institutional\s+mechanism", re.I)),
    ("Caution-Advisory", re.compile(r"\bcaution\w*\b|unsolicited\s+(messages|videos)|\badvisory\b", re.I)),
    ("Vendor-Tech-Admin", re.compile(
        r"vendor\s+empanelment|non.?neat\s+front\s*end|\bNNF\b|member\s+interface|\bUDiFF\b|"
        r"unique\s+(device\s+)?identifier", re.I)),
    ("Client-Due-Diligence", re.compile(r"client\s+due\s+diligence|modification\s+of\s+client\s+codes", re.I)),
    ("Member-Dashboard", re.compile(r"member\s+surveillance\s+dashboard", re.I)),
    ("Derivatives-Contract-Specific", re.compile(
        r"futures\s+and\s+options\s+contracts?\s+in|derivative\s+contracts?\s+in|exclusion\s+of\s+futures", re.I)),
    ("Illiquid-Securities", re.compile(r"illiquid\s+securities", re.I)),
    ("Risk-Controls-Misc", re.compile(
        r"pre-?trade\s+risk|penalty\s+on\s+abnormal|placing\s+(?:of\s+)?orders?\s+at|monitoring\s+of\s+foreign\s+investment|"
        r"operational\s+efficiency\s+in\s+monitoring|market\s+making\s+compliance|"
        r"empanelment\s+of\s+vendors\s+for\s+surveillance\s+software|discontinuation\s+of\s+disclosure|"
        r"extension\s+of\s+time\s+for\s+reporting|electronic\s+gold\s+receipts|revision\s+in\s+dynamic\s+price\s+band",
        re.I)),
)

def is_periodic_asm_subject(subject: str) -> bool:
    """Attempt every SURV circular subject UNLESS it matches a known-irrelevant category in
    `_DENY_PATTERNS` -- see that constant's docstring (P4-013) for why this is a denylist, not an
    allowlist. An empty/missing subject is denied (nothing to attempt)."""
    s = (subject or "").strip()
    if not s:
        return False
    return not any(pat.search(s) for _, pat in _DENY_PATTERNS)

# ---------- Section-aware Annexure parsing (pure, no I/O) ----------

def _cell(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]

def _is_blank(row: tuple) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in row)

def _is_title_row(row: tuple) -> bool:
    first = _cell(row, 0)
    return isinstance(first, str) and bool(TITLE_START_RE.match(first))

def _is_header_row(row: tuple) -> bool:
    return any(isinstance(c, str) and re.sub(r"[.\s]", "", c.lower()) == "srno" for c in row)

def _find_column(header_row: tuple, name: str) -> int | None:
    for idx, c in enumerate(header_row):
        if isinstance(c, str) and c.strip().lower() == name:
            return idx
    return None

def _valid_symbol(value) -> bool:
    return isinstance(value, str) and bool(value.strip()) and value.strip().lower() != "nil"

# '#' observed on real symbols (e.g. "IMAGICAA #", footnoted "# As per BSE") in addition to the
# '*'/'^' already seen on symbols and section titles (P4-010/P4-011, docs/DEFECT_REGISTER.md).
MARKER_CHARS = ("*", "^", "#")

def _footnote_in_row(row: tuple) -> tuple[str, str] | None:
    """A footnote-definition row has exactly one non-blank cell, and that cell starts with a
    marker character -- true regardless of WHICH column it lands in. Real circulars put it in
    column 0 ("* Moved from STASM to LTASM framework", ...) in most sheets, but at least one
    (SURV49783's Annexure I-A) puts it in the Symbol column position instead
    (None, '* Moved from STASM to LTASM framework', None, None) -- a column-0-only check missed
    this, and the footnote text was stored as if it were a real symbol (P4-010)."""
    non_blank = [c.strip() for c in row if isinstance(c, str) and c.strip()]
    if len(non_blank) != 1:
        return None
    text = non_blank[0]
    if text[:1] in MARKER_CHARS:
        return text[0], text[1:].strip()
    return None

def _strip_marker(value):
    """Returns (clean_value, marker_char_or_None). Real symbols carry a trailing marker directly
    ("MARATHON *", "IMAGICAA #") that must be stripped before the value is used as a business
    key -- otherwise the same real company splits into two different stored symbol identities
    depending on which circular did or didn't carry the marker (P4-011,
    docs/DEFECT_REGISTER.md), breaking lifecycle continuity."""
    if not isinstance(value, str):
        return value, None
    v = value.strip()
    if v and v[-1] in MARKER_CHARS:
        return v[:-1].strip(), v[-1]
    return v, None

def _classify_title(title: str) -> tuple[str, str, str | None, str | None] | None:
    """Returns (mechanism, action_type, from_stage, to_stage) or None if `title` isn't one of the
    known ASM stage-event section types (e.g. an IBC carve-out section)."""
    m = TRANSITION_TITLE_RE.search(title)
    if m:
        term, from_stage, to_stage = m.groups()
        return (ASM_LT if term.lower() == "long" else ASM_ST, STAGE_CHANGE, from_stage.upper(), to_stage.upper())
    m = ENTRY_TITLE_RE.search(title)
    if m:
        term, to_stage = m.groups()
        return (ASM_LT if term.lower() == "long" else ASM_ST, ENTRY, None, to_stage.upper())
    m = EXCLUSION_TITLE_RE.search(title)
    if m:
        term = m.group(1)
        mechanism = (ASM_LT if term.lower() == "long" else ASM_ST) if term else None
        return (mechanism, EXIT, None, None)  # None mechanism resolved by caller from the circular subject
    return None

def parse_sectioned_annexure(rows: list[tuple], subject_mechanism: str, source_circular: str,
                              default_knowledge_date: str) -> tuple[list[dict], list[str]]:
    """Pure function: splits one Annexure sheet's raw rows (as `ws.iter_rows(values_only=True)`
    would yield) into its titled sub-sections and emits one event dict per real symbol row.
    `subject_mechanism` (ASM_LT/ASM_ST, from the circular's own subject line) is used only for
    exclusion sections, whose title carries no Long/Short wording of its own.
    Returns (events, unparsed_titles) -- a title this function doesn't recognize is never silently
    dropped, it comes back in `unparsed_titles` for the caller to report.
    """
    events: list[dict] = []
    unparsed_titles: list[str] = []
    i, n = 0, len(rows)

    while i < n:
        row = rows[i]
        if not _is_title_row(row):
            i += 1
            continue

        title = _cell(row, 0).strip()
        wef_m = WEF_DATE_RE.search(title)
        event_date = to_iso_date(wef_m.group(1)) if wef_m else default_knowledge_date
        section_marker = None
        marker_m = re.search(r"([\^*])\s*w\.e\.f", title, re.I)
        if marker_m:
            section_marker = marker_m.group(1)

        i += 1
        while i < n and _is_blank(rows[i]):
            i += 1
        if i >= n or not _is_header_row(rows[i]):
            unparsed_titles.append(title)
            continue
        header = rows[i]
        symbol_idx = _find_column(header, "symbol")
        name_idx = _find_column(header, "security name")
        i += 1

        data_rows: list[tuple] = []
        footnotes: dict[str, str] = {}
        while i < n and not _is_title_row(rows[i]):
            r = rows[i]
            if _is_blank(r):
                i += 1
                continue
            fn = _footnote_in_row(r)
            if fn is not None:
                marker, text = fn
                footnotes[marker] = text
                i += 1
                continue
            data_rows.append(r)
            i += 1

        classification = _classify_title(title)
        if classification is None:
            unparsed_titles.append(title)
            continue
        mechanism, action_type, from_stage, to_stage = classification
        if mechanism is None:
            mechanism = subject_mechanism

        for r in data_rows:
            raw_symbol = _cell(r, symbol_idx)
            symbol, symbol_marker = _strip_marker(raw_symbol)
            if not _valid_symbol(symbol):
                continue
            raw_name = _cell(r, name_idx) if name_idx is not None else None
            _, name_marker = _strip_marker(raw_name)
            row_marker = symbol_marker or name_marker
            if row_marker:
                details = footnotes.get(row_marker)
            elif section_marker:
                details = footnotes.get(section_marker)
            else:
                details = None
            events.append({
                "symbol": symbol.strip(), "mechanism": mechanism, "action_type": action_type,
                "from_stage": from_stage, "to_stage": to_stage,
                "event_date": event_date, "knowledge_date": default_knowledge_date,
                "source_circular": source_circular, "details": details,
            })

    return events, unparsed_titles

def parse_circular_workbook(wb, subject: str, source_circular: str, knowledge_date: str) -> tuple[list[dict], list[str]]:
    """Parses every `Annexure*` sheet in an opened workbook (openpyxl `Workbook`). Deliberately
    skips `Consolidated*` sheets -- see module docstring."""
    subject_mechanism = mechanism_from_subject(subject)
    all_events: list[dict] = []
    all_unparsed: list[str] = []
    for sheet_name in wb.sheetnames:
        if not sheet_name.strip().lower().startswith("annexure"):
            continue
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        events, unparsed = parse_sectioned_annexure(rows, subject_mechanism, source_circular, knowledge_date)
        all_events.extend(events)
        all_unparsed.extend(unparsed)
    return all_events, all_unparsed

def parse_circular_zip_bytes(zip_bytes: bytes, subject: str, source_circular: str, knowledge_date: str) -> tuple[list[dict], list[str]]:
    with zipfile.ZipFile(BytesIO(zip_bytes)) as zf:
        # Match on the basename, not the full zip-entry path: some real circulars (e.g. SURV64061)
        # nest the workbook inside a subfolder ("SURV64061/Annexure.xlsx"), which a full-path
        # `startswith("annexure")` check misses entirely (P4-007, docs/DEFECT_REGISTER.md).
        xlsx_names = [nm for nm in zf.namelist()
                      if nm.replace("\\", "/").rsplit("/", 1)[-1].lower().startswith("annexure")
                      and nm.lower().endswith(".xlsx")]
        if not xlsx_names:
            raise ValueError(f"No Annexure*.xlsx found in zip; contents={zf.namelist()}")
        xlsx_bytes = zf.read(xlsx_names[0])
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True)
    return parse_circular_workbook(wb, subject, source_circular, knowledge_date)

# ---------- Bulk ingestion + reporting ----------

@dataclass
class AsmIngestionReport:
    event_counts: dict[str, int] = field(default_factory=dict)  # PARSED counts, see duplicate_events_skipped
    duplicate_events_skipped: int = 0
    circulars_processed: int = 0
    circulars_skipped_not_periodic: int = 0
    circulars_failed: list[dict] = field(default_factory=list)  # {"circular", "reason"}
    unparsed_section_titles: list[dict] = field(default_factory=list)  # {"circular", "title"}

def ingest_asm_circular(conn, subject: str, source_circular: str, knowledge_date: str,
                         zip_bytes: bytes, report: AsmIngestionReport) -> None:
    """Parses and writes one circular's events through the Phase 1 store. Raises on a
    `StoreValidationError` (our own contract violation -- must never be silently logged as a
    source-format failure); any other parse exception is caught, counted, and logged by the
    caller's sweep loop, not raised, since a single malformed real-world circular must not abort
    the whole ingestion run.

    `event_counts` reflects PARSED events, not necessarily inserted ones: two different circulars
    can legitimately describe the identical fact (same symbol/mechanism/event_date/knowledge_date
    -- a same-day corrective/reissue circular pair is a real, observed NSE pattern, confirmed for
    SURV68086/SURV68088 both naming ICDSLTD), and `write_facts` correctly skips the second as a
    duplicate rather than raising. `duplicate_events_skipped` (from `write_facts`'s own
    `BulkWriteResult.skipped_duplicate`, P4-009 -- docs/DEFECT_REGISTER.md) is what makes
    `sum(event_counts) - duplicate_events_skipped == COUNT(*) in the store` hold exactly; treating
    `event_counts` alone as "rows written" is exactly the P4-004/P4-005 mistake one level up.
    """
    events, unparsed = parse_circular_zip_bytes(zip_bytes, subject, source_circular, knowledge_date)
    if events:
        # Must come BEFORE circulars_processed is incremented (P4-005, docs/DEFECT_REGISTER.md):
        # if this write raises, the caller's sweep loop counts the circular as failed, and it must
        # not ALSO be counted as processed just because parsing got that far.
        result = write_facts(conn, "surveillance_flags", [{**e, "source_file": source_circular} for e in events])
        report.duplicate_events_skipped += result.skipped_duplicate
    report.circulars_processed += 1
    for title in unparsed:
        report.unparsed_section_titles.append({"circular": source_circular, "title": title})
    for e in events:
        key = f'{e["mechanism"]}:{e["action_type"]}'
        report.event_counts[key] = report.event_counts.get(key, 0) + 1

# ---------- Real network fetch (not used by the fixture-based test suite) ----------

def session_with_cookie(timeout: float = 20.0) -> requests.Session:
    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36",
        "Accept": "application/json",
    })
    session.get("https://www.nseindia.com", timeout=timeout)
    return session

def fetch_circular_index(session: requests.Session, from_date: date, to_date: date, timeout: float = 30.0) -> list[dict]:
    """P8-004 (docs/DEFECT_REGISTER.md): the real, live response from this endpoint is an envelope
    -- {"data": [...circulars...], "fromDate": ..., "toDate": ...} -- not a bare list, confirmed
    directly against the real API (26 real circulars returned for a real 10-day window, 2026-09).
    This function was never exercised against live data anywhere in this project's test suite
    before that check (`fetch_and_ingest_asm_range`'s own docstring: "Not called by the
    fixture-based test suite") -- `for c in circulars` previously iterated the envelope dict's own
    KEYS ("data"/"fromDate"/"toDate", each a string), not its circulars, crashing on the first
    real call with `AttributeError: 'str' object has no attribute 'get'`."""
    r = session.get("https://www.nseindia.com/api/circulars",
                     params={"dept": "SURV", "fromDate": from_date.strftime("%d-%m-%Y"), "toDate": to_date.strftime("%d-%m-%Y")},
                     timeout=timeout)
    r.raise_for_status()
    return r.json()["data"]

def fetch_circular_file(session: requests.Session, url: str, timeout: float = 30.0) -> bytes:
    r = session.get(url, timeout=timeout)
    r.raise_for_status()
    return r.content

def fetch_and_ingest_asm_range(conn, from_date: date, to_date: date, delay_seconds: float = 0.3,
                                session_factory: Callable[[], requests.Session] = session_with_cookie) -> AsmIngestionReport:
    """Real network sweep: lists every SURV circular in the date range, keeps only periodic ASM
    applicability circulars (`is_periodic_asm_subject`), downloads+parses+writes each one through
    the store. Not called by the fixture-based test suite."""
    session = session_factory()
    report = AsmIngestionReport()
    circulars = fetch_circular_index(session, from_date, to_date)
    for c in circulars:
        subject = c.get("sub", "") or ""
        if not is_periodic_asm_subject(subject):
            report.circulars_skipped_not_periodic += 1
            continue
        source_circular = "SURV" + str(c["circNumber"])
        knowledge_date = f'{c["cirDate"][:4]}-{c["cirDate"][4:6]}-{c["cirDate"][6:8]}'
        try:
            zip_bytes = fetch_circular_file(session, c["circFilelink"])
            ingest_asm_circular(conn, subject, source_circular, knowledge_date, zip_bytes, report)
        except StoreValidationError:
            raise
        except Exception as exc:
            report.circulars_failed.append({"circular": source_circular, "reason": str(exc)})
        time.sleep(delay_seconds)
    return report
